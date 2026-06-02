"""
process_event_samples.py
========================
Combines CSV exports from Azure (WindowsEvent) and XSIAM (microsoft_windows_raw)
into a single formatted Excel workbook for the Cribl Stream parsing team.

Usage:
    python process_event_samples.py \
        --azure  azure_windowsevent_samples.csv \
        --xsiam  xsiam_windows_raw_samples.csv \
        --output event_samples_for_parsing.xlsx

Requirements:
    pip install pandas openpyxl
"""

import argparse
import json
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Helpers ──────────────────────────────────────────────────────────────────

def load_csv(path: str, label: str) -> pd.DataFrame:
    try:
        df = pd.read_csv(path, dtype=str).fillna("")
        print(f"[{label}] Loaded {len(df)} rows from {path}")
        return df
    except FileNotFoundError:
        print(f"[WARN] {path} not found — skipping {label} sheet")
        return pd.DataFrame()


def normalise_event_id(df: pd.DataFrame, candidates: list[str]) -> pd.DataFrame:
    """Find the EventID column regardless of capitalisation/naming."""
    for col in candidates:
        if col in df.columns:
            df = df.rename(columns={col: "EventID"})
            df["EventID"] = pd.to_numeric(df["EventID"], errors="coerce").astype("Int64")
            return df.sort_values("EventID")
    print(f"[WARN] Could not find EventID column. Tried: {candidates}")
    return df


def style_sheet(ws, header_color: str = "1F4E79"):
    """Apply header styling and auto-width to a worksheet."""
    header_fill = PatternFill("solid", fgColor=header_color)
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def truncate_eventdata(df: pd.DataFrame, col: str, max_len: int = 500) -> pd.DataFrame:
    """Truncate long EventData strings so Excel cells don't overflow."""
    if col in df.columns:
        df[col] = df[col].apply(
            lambda v: (v[:max_len] + "…") if isinstance(v, str) and len(v) > max_len else v
        )
    return df


# ── Main ─────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Build event-sample Excel from Azure + XSIAM CSVs")
    parser.add_argument("--azure",  default="azure_windowsevent_samples.csv",  help="Azure CSV export")
    parser.add_argument("--xsiam",  default="xsiam_windows_raw_samples.csv",   help="XSIAM CSV export")
    parser.add_argument("--output", default="event_samples_for_parsing.xlsx",  help="Output Excel file")
    args = parser.parse_args()

    azure_df = load_csv(args.azure, "Azure")
    xsiam_df = load_csv(args.xsiam, "XSIAM")

    # Normalise EventID columns
    if not azure_df.empty:
        azure_df = normalise_event_id(azure_df, ["EventID", "eventid", "event_id"])
        azure_df = truncate_eventdata(azure_df, "EventData")
        azure_df = truncate_eventdata(azure_df, "RawEventData")

    if not xsiam_df.empty:
        xsiam_df = normalise_event_id(xsiam_df, ["event_id", "EventID", "eventid"])
        xsiam_df = truncate_eventdata(xsiam_df, "action_evtlog_data")
        xsiam_df = truncate_eventdata(xsiam_df, "raw_log")

    # ── Comparison sheet ────────────────────────────────────────────────────
    comparison_rows = []
    azure_ids = set(azure_df["EventID"].dropna().unique()) if not azure_df.empty and "EventID" in azure_df.columns else set()
    xsiam_ids = set(xsiam_df["EventID"].dropna().unique()) if not xsiam_df.empty and "EventID" in xsiam_df.columns else set()
    all_ids   = sorted(azure_ids | xsiam_ids)

    for eid in all_ids:
        in_azure = eid in azure_ids
        in_xsiam = eid in xsiam_ids
        status = "Both" if (in_azure and in_xsiam) else ("Azure only" if in_azure else "XSIAM only")
        comparison_rows.append({
            "EventID":    eid,
            "In Azure":   "✓" if in_azure else "",
            "In XSIAM":   "✓" if in_xsiam else "",
            "Status":     status,
            "Notes":      "",   # leave blank for project team to fill
        })

    comparison_df = pd.DataFrame(comparison_rows)

    print(f"\nSummary:")
    print(f"  Azure unique EventIDs : {len(azure_ids)}")
    print(f"  XSIAM unique EventIDs : {len(xsiam_ids)}")
    print(f"  In both environments  : {len(azure_ids & xsiam_ids)}")
    print(f"  Azure only            : {len(azure_ids - xsiam_ids)}")
    print(f"  XSIAM only            : {len(xsiam_ids - azure_ids)}")

    # ── Write Excel ─────────────────────────────────────────────────────────
    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        comparison_df.to_excel(writer, sheet_name="EventID Comparison", index=False)
        if not azure_df.empty:
            azure_df.to_excel(writer, sheet_name="Azure - WindowsEvent", index=False)
        if not xsiam_df.empty:
            xsiam_df.to_excel(writer, sheet_name="XSIAM - windows_raw", index=False)

    # Apply styling
    wb = load_workbook(args.output)
    style_sheet(wb["EventID Comparison"], header_color="1F4E79")
    if "Azure - WindowsEvent" in wb.sheetnames:
        style_sheet(wb["Azure - WindowsEvent"], header_color="0070C0")
    if "XSIAM - windows_raw" in wb.sheetnames:
        style_sheet(wb["XSIAM - windows_raw"], header_color="375623")
    wb.save(args.output)

    print(f"\n✓ Excel written to: {args.output}")


if __name__ == "__main__":
    main()
